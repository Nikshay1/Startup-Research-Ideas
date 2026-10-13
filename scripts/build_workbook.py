#!/usr/bin/env python3
"""Build output/startup-idea-research.xlsx from the CSV files in data/.

The CSVs are the source of truth. This script only lays them out:
one tab per theme (bracket), plus overview, incubator and raw YC tabs.

Usage:  uv run --with openpyxl python3 scripts/build_workbook.py
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output" / "startup-idea-research.xlsx"
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INDIA_FILL = PatternFill("solid", fgColor="FFF2CC")
BATCHES = ["Fall 2025", "Winter 2026", "Spring 2026", "Summer 2026", "Fall 2026"]
YC_SHEET = "YC all 900"

# (column in funded_startups.csv, header shown in the sheet, width)
STARTUP_COLS = [
    ("idea_simple", "Idea in simple words", 70),
    ("startup", "Startup", 24),
    ("funder", "Funded by", 26),
    ("program_batch", "Batch / cohort", 24),
    ("funding_amount", "Funding", 18),
    ("funding_note", "What the funding number means", 46),
    ("based_in", "Based in", 28),
    ("region", "India / USA / Other", 14),
    ("india_link", "Indian founders?", 14),
    ("website", "Website", 28),
    ("source_url", "Source", 40),
    ("confidence", "Confidence", 12),
]


def read(name):
    with open(ROOT / "data" / name, newline="") as f:
        return list(csv.DictReader(f))


def write_table(ws, headers, widths, rows, start_row=1, wrap_cols=(), india_col=None):
    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=start_row, column=col, value=head)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top", wrap_text=(col - 1) in wrap_cols)
        if india_col is not None and row[india_col] == "India":
            for col in range(1, len(row) + 1):
                ws.cell(row=r, column=col).fill = INDIA_FILL
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{start_row}:{last}{start_row + len(rows)}"


def startup_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    keys = [k for k, _, _ in STARTUP_COLS]
    write_table(ws, [h for _, h, _ in STARTUP_COLS], [w for _, _, w in STARTUP_COLS],
                [[r[k] for k in keys] for r in rows], wrap_cols=(0, 5), india_col=keys.index("region"))
    return ws


def main():
    brackets = read("brackets.csv")
    startups = read("funded_startups.csv")
    incubators = read("incubators.csv")
    yc = read("yc_recent_batches.csv")

    wb = Workbook()

    # --- Start here -----------------------------------------------------
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Startup idea research", "Which pre-seed ideas are incubators funding, and where can two founders from India apply?"),
        ("", ""),
        ("How to read this file", ""),
        ("Themes overview", "All themes on one page with how many companies YC funded in each of its last 5 batches. Start here to see what is rising."),
        ("Incubators", "Each program: money, equity taken, who can apply, timing, and how far to trust the row."),
        ("One tab per theme", "Each theme tab lists funded startups with the idea in one plain sentence. Rows based in India are shaded yellow."),
        ("India-based", "Every funded startup in this file that is based in India, across all themes."),
        (YC_SHEET, "Every YC company from Fall 2025 to Fall 2026 with YC's own one-line description and its theme. Use the filters."),
        ("", ""),
        ("What the columns mean", ""),
        ("Funding", "The program's standard cheque unless the startup's own round was reported. The next column says which."),
        ("India / USA / Other", "Where the company is based, not where the founders are from. 'Unknown' means the source did not say."),
        ("Indian founders?", "yes = source says so or company is in India. reported = a newsletter lists them as Indian-origin, not verified. unknown = not checked."),
        ("Confidence", "high = official directory or the program's own page. medium = news report. low = third-party list."),
        ("", ""),
        ("Where the data lives", "This file is generated from the CSV files in data/ by scripts/build_workbook.py. Edit the CSVs, not this file."),
        ("Research date", "2026-10-13"),
    ]
    for r, (a, b) in enumerate(lines, start=1):
        ca, cb = ws.cell(row=r, column=1, value=a), ws.cell(row=r, column=2, value=b)
        ca.font = Font(name=FONT, bold=True, size=14 if r == 1 else 11)
        cb.font = Font(name=FONT)
        cb.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Themes overview --------------------------------------------------
    ws = wb.create_sheet("Themes overview")
    heads = ["Theme", "What it means in simple words", "Tab name", "Startups listed in this file"]
    heads += [f"YC {b}" for b in BATCHES] + ["YC total (5 batches)", "Share of Fall 2025 batch", "Share of Summer 2026 batch"]
    rows = []
    first = 2
    last = first + len(brackets) - 1
    yc_ref = f"'{YC_SHEET}'"
    for i, b in enumerate(brackets):
        r = first + i
        counts = [f'=COUNTIFS({yc_ref}!$A:$A,$C{r},{yc_ref}!$E:$E,"{batch}")' for batch in BATCHES]
        rows.append([b["bracket_name"], b["what_it_means_in_simple_words"], b["sheet_name"],
                     f"=COUNTA('{b['sheet_name']}'!$B:$B)-1", *counts, f"=SUM(E{r}:I{r})",
                     f"=IF(E${last + 1}=0,0,E{r}/E${last + 1})", f"=IF(H${last + 1}=0,0,H{r}/H${last + 1})"])
    write_table(ws, heads, [34, 80, 22, 14, 11, 11, 11, 11, 11, 13, 13, 13], rows, wrap_cols=(1,))
    total_row = last + 1
    ws.cell(row=total_row, column=1, value="Total").font = Font(name=FONT, bold=True)
    for col in range(4, 11):
        letter = get_column_letter(col)
        c = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
        c.font = Font(name=FONT, bold=True)
    for r in range(first, last + 1):
        ws.cell(row=r, column=11).number_format = "0.0%"
        ws.cell(row=r, column=12).number_format = "0.0%"
    note = ws.cell(row=total_row + 2, column=1, value=(
        "YC counts are formulas over the 'YC all 900' tab (source: yc-oss.github.io/api, read 2026-10-13). "
        "Themes were assigned by hand from each company's description; see data/yc_bracket_assignments.csv."))
    note.font = Font(name=FONT, italic=True)

    # --- Incubators ---------------------------------------------------------
    ws = wb.create_sheet("Incubators")
    inc_cols = [("program", "Program", 30), ("type", "Type", 18), ("based_in", "Based in", 22),
                ("open_to_founders_in_india", "Can founders in India apply?", 40), ("money", "Money", 26),
                ("equity_taken", "Equity taken", 40), ("stage_and_who", "Stage and who they want", 44),
                ("format", "Format", 40), ("student_friendly", "Student friendly?", 34),
                ("timing", "Timing / next date", 40), ("notes", "Notes", 50),
                ("source_url", "Source", 40), ("confidence", "Confidence", 12)]
    write_table(ws, [h for _, h, _ in inc_cols], [w for _, _, w in inc_cols],
                [[r[k] for k, _, _ in inc_cols] for r in incubators], wrap_cols=tuple(range(3, 11)))

    # --- One tab per theme ----------------------------------------------------
    for b in brackets:
        startup_sheet(wb, b["sheet_name"], [s for s in startups if s["bracket_id"] == b["bracket_id"]])

    # --- India-based ------------------------------------------------------------
    startup_sheet(wb, "India-based", [s for s in startups if s["region"] == "India"])

    # --- Raw YC -----------------------------------------------------------------
    ws = wb.create_sheet(YC_SHEET)
    names = {b["bracket_id"]: b["sheet_name"] for b in brackets}
    yc_cols = [("Theme (tab name)", 22), ("Startup", 26), ("YC's one-line description", 60), ("Full description", 80),
               ("Batch", 14), ("India / USA / Other", 14), ("Location", 28), ("Team size", 10),
               ("YC industry", 34), ("Website", 30), ("YC page", 44)]
    rows = [[names[c["bracket_id"]], c["startup"], c["yc_one_liner"], c["yc_long_description"], c["batch"],
             c["region"], c["location"], int(c["team_size"]) if c["team_size"] else None, c["yc_industry"],
             c["website"], c["yc_url"]] for c in yc]
    write_table(ws, [h for h, _ in yc_cols], [w for _, w in yc_cols], rows, india_col=5)

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} with {len(wb.sheetnames)} tabs")


if __name__ == "__main__":
    main()
